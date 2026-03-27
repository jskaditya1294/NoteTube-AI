# from __future__ import annotations
# from datetime import datetime
# from pathlib import Path
# import re
# from typing import Optional, Dict
# from docx import Document


# def _safe_filename(value: str) -> str:
#     """Make a string safe for Windows file names."""
#     value = value.strip()
#     value = re.sub(r'[<>:"/\\|?*\n]+', "_", value)
#     value = re.sub(r"\s+", "_", value)
#     return value[:120] if value else "notes"


# def _build_notes_markdown(
#     notes: str,
#     video_id: Optional[str],
#     title: Optional[str],
#     created_at: str,
# ) -> str:
#     heading = title.strip() if title else f"YouTube Notes ({video_id or 'unknown'})"

#     lines = [
#         f"# {heading}",
#         "",
#         "## Metadata",
#         f"- **Created at:** {created_at}",
#         f"- **Video ID:** {video_id or 'N/A'}",
#         "",
#         "## Notes",
#         "",
#     ]

#     # Keep the notes content as-is so markdown headings/bullets survive
#     lines.append((notes or "").strip())
#     lines.append("")

#     return "\n".join(lines).strip() + "\n"


# def _write_notes_to_docx(
#     notes: str,
#     doc: Document,
# ) -> None:
#     """
#     Preserve simple markdown-like structure in DOCX:
#     # Heading      -> heading level 1
#     ## Heading     -> heading level 2
#     - item         -> bullet list
#     plain text     -> paragraph
#     """
#     for raw_line in (notes or "").splitlines():
#         line = raw_line.strip()

#         if not line:
#             doc.add_paragraph("")
#             continue

#         if line.startswith("# "):
#             doc.add_heading(line[2:].strip(), level=1)
#         elif line.startswith("## "):
#             doc.add_heading(line[3:].strip(), level=2)
#         elif line.startswith("### "):
#             doc.add_heading(line[4:].strip(), level=3)
#         elif line.startswith("- "):
#             doc.add_paragraph(line[2:].strip(), style="List Bullet")
#         else:
#             doc.add_paragraph(line)


# def save_notes_as_md_and_docx(
#     notes: str,
#     output_dir: str | Path,
#     video_id: Optional[str] = None,
#     title: Optional[str] = None,
# ) -> Dict[str, str]:
#     """
#     Save final notes in both Markdown and Word formats.

#     Returns:
#         dict with keys: markdown_path, docx_path
#     """
#     if not notes or not notes.strip():
#         raise ValueError("Notes are empty; nothing to save.")

#     out_dir = Path(output_dir)
#     out_dir.mkdir(parents=True, exist_ok=True)

#     created_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
#     base = _safe_filename(title or video_id or "notes")

#     md_path = out_dir / f"{base}_notes.md"
#     docx_path = out_dir / f"{base}_notes.docx"

#     # Markdown
#     markdown_content = _build_notes_markdown(
#         notes=notes,
#         video_id=video_id,
#         title=title,
#         created_at=created_at,
#     )
#     md_path.write_text(markdown_content, encoding="utf-8")

#     # Word
#     doc = Document()
#     doc.add_heading(title or f"YouTube Notes ({video_id or 'unknown'})", level=1)
#     doc.add_heading("Metadata", level=2)
#     doc.add_paragraph(f"Created at: {created_at}")
#     doc.add_paragraph(f"Video ID: {video_id or 'N/A'}")
#     doc.add_heading("Notes", level=2)

#     _write_notes_to_docx(notes, doc)

#     doc.save(docx_path)

#     return {
#         "markdown_path": str(md_path.resolve()),
#         "docx_path": str(docx_path.resolve()),
#     }



from __future__ import annotations
from datetime import datetime
from pathlib import Path
import re
from typing import Optional, Dict, List

from docx import Document
from docx.shared import Inches


def _safe_filename(value: str) -> str:
    """Make a string safe for Windows file names."""
    value = value.strip()
    value = re.sub(r'[<>:"/\\|?*\n]+', "_", value)
    value = re.sub(r"\s+", "_", value)
    return value[:120] if value else "notes"


def _build_notes_markdown(
    notes: str,
    video_id: Optional[str],
    title: Optional[str],
    created_at: str,
    important_frames: Optional[List[dict]] = None,
    output_dir: Optional[Path] = None,
) -> str:
    heading = title.strip() if title else f"YouTube Notes ({video_id or 'unknown'})"

    lines = [
        f"# {heading}",
        "",
        "## Metadata",
        f"- **Created at:** {created_at}",
        f"- **Video ID:** {video_id or 'N/A'}",
        "",
        "## Notes",
        "",
    ]
    lines.append((notes or "").strip())
    lines.append("")

    if important_frames and output_dir is not None:
        out = Path(output_dir)
        lines.append("## Key images")
        lines.append("")
        for frame in important_frames:
            path = frame.get("path")
            caption = frame.get("caption") or "Key frame"
            ts = frame.get("timestamp_sec")
            if path and Path(path).exists():
                try:
                    rel = Path(path).relative_to(out)
                    lines.append(f"### {caption}")
                    if ts is not None:
                        lines.append(f"*Timestamp: {int(ts)}s*")
                    lines.append(f"![{caption}]({rel.as_posix()})")
                    lines.append("")
                except ValueError:
                    lines.append(f"### {caption}")
                    if ts is not None:
                        lines.append(f"*Timestamp: {int(ts)}s*")
                    lines.append("")

    return "\n".join(lines).strip() + "\n"


def _write_notes_to_docx(
    notes: str,
    doc: Document,
) -> None:
    """
    Preserve markdown-like structure in DOCX:
    - # Heading      -> heading level 1
    - ## Heading     -> heading level 2
    - ### Heading    -> heading level 3
    - #### Heading   -> heading level 4
    - -, *, • item   -> bullet list (nested by indent)
    - plain text     -> paragraph

    Rules:
    - Heading is recognized only when it starts at column 0.
    - Bullet nesting is derived from leading spaces before stripping.
    """
    indent_step = 2
    max_bullet_level = 2  # 0, 1, 2 -> up to 3 Word bullet levels

    bullet_styles = {
        0: "List Bullet",
        1: "List Bullet 2",
        2: "List Bullet 3",
    }

    def _add_bullet_paragraph(text: str, level: int) -> None:
        """Add a bullet paragraph with style by level; fallback to manual indent if needed."""
        level = min(level, max_bullet_level)
        style_name = bullet_styles[level]

        try:
            doc.add_paragraph(text, style=style_name)
        except KeyError:
            # Fallback if template/style is missing
            p = doc.add_paragraph(text, style="List Bullet")
            p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
            p.paragraph_format.first_line_indent = Inches(-0.18)

    for raw_line in (notes or "").splitlines():
        # Expand tabs so indentation behaves predictably
        expanded_line = raw_line.expandtabs(indent_step)
        stripped = expanded_line.strip()

        # Empty line
        if not stripped:
            doc.add_paragraph("")
            continue

        # Preserve indent before stripping
        leading_spaces = len(expanded_line) - len(expanded_line.lstrip(" "))
        bullet_level = min(leading_spaces // indent_step, max_bullet_level)

        # Headings only when line starts at column 0
        if leading_spaces == 0 and stripped.startswith("#### "):
            doc.add_heading(stripped[5:].strip(), level=4)
        elif leading_spaces == 0 and stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=3)
        elif leading_spaces == 0 and stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=2)
        elif leading_spaces == 0 and stripped.startswith("# "):
            doc.add_heading(stripped[2:].strip(), level=1)

        # Bullets at any indent
        elif stripped.startswith("- ") or stripped.startswith("* "):
            content = stripped[2:].strip()
            _add_bullet_paragraph(content, bullet_level)
        elif stripped.startswith("• "):
            content = stripped[2:].strip()
            _add_bullet_paragraph(content, bullet_level)

        # Normal paragraph
        else:
            doc.add_paragraph(stripped)


def save_notes_as_md_and_docx(
    notes: str,
    output_dir: str | Path,
    video_id: Optional[str] = None,
    title: Optional[str] = None,
    important_frames: Optional[List[dict]] = None,
) -> Dict[str, str]:
    """
    Save final notes in both Markdown and Word formats.
    If important_frames is provided, adds a "Key images" section.
    """
    if not notes or not notes.strip():
        raise ValueError("Notes are empty; nothing to save.")

    out_dir = Path(output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    created_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    base = _safe_filename(title or video_id or "notes")

    md_path = out_dir / f"{base}_notes.md"
    docx_path = out_dir / f"{base}_notes.docx"

    markdown_content = _build_notes_markdown(
        notes=notes,
        video_id=video_id,
        title=title,
        created_at=created_at,
        important_frames=important_frames or [],
        output_dir=out_dir,
    )
    md_path.write_text(markdown_content, encoding="utf-8")

    doc = Document()
    doc.add_heading(title or f"YouTube Notes ({video_id or 'unknown'})", level=1)
    doc.add_heading("Metadata", level=2)
    doc.add_paragraph(f"Created at: {created_at}")
    doc.add_paragraph(f"Video ID: {video_id or 'N/A'}")
    doc.add_heading("Notes", level=2)
    _write_notes_to_docx(notes, doc)

    if important_frames:
        doc.add_heading("Key images", level=2)
        for frame in important_frames:
            path = frame.get("path")
            caption = frame.get("caption") or "Key frame"
            ts = frame.get("timestamp_sec")
            if path and Path(path).exists():
                try:
                    doc.add_paragraph(f"{caption}" + (f" (at {int(ts)}s)" if ts is not None else ""))
                    doc.add_picture(path, width=Inches(5.5))
                    doc.add_paragraph("")
                except Exception:
                    doc.add_paragraph(f"{caption} — [Image: {path}]")

    doc.save(docx_path)
    return {"markdown_path": str(md_path.resolve()), "docx_path": str(docx_path.resolve())}


def _write_interview_markdown_to_docx(doc: Document, section_md: str) -> None:
    """Append interview bank markdown-ish content to a Word document."""
    for raw in section_md.splitlines():
        line = raw.rstrip()
        stripped = line.strip()
        if not stripped:
            doc.add_paragraph("")
        elif stripped.startswith("#### "):
            doc.add_heading(stripped[5:].strip(), level=4)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=2)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:].strip(), level=1)
        elif stripped.startswith("|") and "|" in stripped[1:]:
            doc.add_paragraph(stripped)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            try:
                doc.add_paragraph(stripped[2:].strip(), style="List Bullet")
            except Exception:
                doc.add_paragraph(stripped[2:].strip())
        elif stripped.startswith("  - ") or stripped.startswith("  * "):
            doc.add_paragraph(stripped.strip())
        else:
            doc.add_paragraph(stripped)


def append_interview_to_notes_files(
    markdown_path: str | Path,
    docx_path: str | Path,
    interview_markdown: str,
) -> None:
    """Append interview question bank to existing notes .md and .docx."""
    if not (interview_markdown or "").strip():
        return
    md_p = Path(markdown_path)
    doc_p = Path(docx_path)
    body = md_p.read_text(encoding="utf-8") if md_p.exists() else ""
    md_p.write_text(body.rstrip() + "\n\n" + interview_markdown.strip() + "\n", encoding="utf-8")
    doc = Document(str(doc_p))
    doc.add_page_break()
    _write_interview_markdown_to_docx(doc, interview_markdown)
    doc.save(str(doc_p))


def save_interview_bank_standalone(
    output_dir: str | Path,
    interview_markdown: str,
) -> Dict[str, str]:
    """Write interview_question_bank.md and interview_question_bank.docx."""
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)
    md_path = out / "interview_question_bank.md"
    docx_path = out / "interview_question_bank.docx"
    md_path.write_text(
        "# Interview question bank\n\n"
        "*MLE / Data Scientist / AI Engineer — sourced from web; not invented.*\n\n"
        + interview_markdown.strip()
        + "\n",
        encoding="utf-8",
    )
    doc = Document()
    doc.add_heading("Interview question bank", level=1)
    doc.add_paragraph("MLE / Data Scientist / AI Engineer — sourced from web; not invented.")
    _write_interview_markdown_to_docx(doc, interview_markdown)
    doc.save(str(docx_path))
    return {"interview_md": str(md_path.resolve()), "interview_docx": str(docx_path.resolve())}


def save_notes_qa_standalone(
    output_dir: str | Path,
    qa_markdown: str,
    title: Optional[str] = None,
) -> Dict[str, str]:
    """Save LLM-generated Q&A as separate .md and .docx files."""
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)

    base = _safe_filename(title or "notes") if title else "notes"
    md_path = out / f"{base}_QA.md"
    docx_path = out / f"{base}_QA.docx"

    heading = title.strip() if title else "Interview Q&A from Notes"
    md_content = (
        f"# {heading} — Interview Q&A\n\n"
        f"*Questions generated from study notes. Answers are based solely on the transcript content.*\n\n"
        + qa_markdown.strip()
        + "\n"
    )
    md_path.write_text(md_content, encoding="utf-8")

    doc = Document()
    doc.add_heading(f"{heading} — Interview Q&A", level=1)
    doc.add_paragraph("Questions generated from study notes. Answers are based solely on the transcript content.")
    _write_interview_markdown_to_docx(doc, qa_markdown)
    doc.save(str(docx_path))

    return {"qa_md": str(md_path.resolve()), "qa_docx": str(docx_path.resolve())}


def save_combined_questions(
    output_dir: str | Path,
    qa_markdown: str,
    interview_markdown: str,
    title: Optional[str] = None,
) -> Dict[str, str]:
    """Save combined Q&A (LLM-generated + web-sourced interview bank) as .md and .docx."""
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)

    base = _safe_filename(title or "notes")
    md_path = out / f"{base}_questions.md"
    docx_path = out / f"{base}_questions.docx"

    heading = title.strip() if title else "Interview Questions"
    lines = [f"# {heading} — Questions & Answers", ""]

    qa_text = (qa_markdown or "").strip()
    interview_text = (interview_markdown or "").strip()

    if qa_text:
        lines.extend([
            "---",
            "",
            "## Part 1: Questions Generated by LLM (from notes)",
            "",
            "*These questions are generated from the study notes using an LLM. Answers are based solely on the transcript content.*",
            "",
            qa_text,
            "",
        ])

    if interview_text:
        lines.extend([
            "---",
            "",
            "## Part 2: Real Interview Questions (sourced from web)",
            "",
            "*These questions were found on the web and attributed to real companies/interviews.*",
            "",
            interview_text,
            "",
        ])

    if not qa_text and not interview_text:
        lines.append("*No questions were generated or found.*")
        lines.append("")

    md_content = "\n".join(lines).strip() + "\n"
    md_path.write_text(md_content, encoding="utf-8")

    doc = Document()
    doc.add_heading(f"{heading} — Questions & Answers", level=1)

    if qa_text:
        doc.add_heading("Part 1: Questions Generated by LLM (from notes)", level=2)
        doc.add_paragraph(
            "These questions are generated from the study notes using an LLM. "
            "Answers are based solely on the transcript content."
        )
        _write_interview_markdown_to_docx(doc, qa_text)

    if interview_text:
        if qa_text:
            doc.add_page_break()
        doc.add_heading("Part 2: Real Interview Questions (sourced from web)", level=2)
        doc.add_paragraph(
            "These questions were found on the web and attributed to real companies/interviews."
        )
        _write_interview_markdown_to_docx(doc, interview_text)

    if not qa_text and not interview_text:
        doc.add_paragraph("No questions were generated or found.")

    doc.save(str(docx_path))
    return {"questions_md": str(md_path.resolve()), "questions_docx": str(docx_path.resolve())}